"""Builds reports/Freight_Rate_Report.docx from the backtest table, the data checks and the December chart.

Run after backtest.py and score.py:  python src/make_report.py
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"


def data_checks():
    tt = pd.read_csv(ROOT / "data" / "train_test.csv", parse_dates=["date"])
    va = pd.read_csv(ROOT / "data" / "validation.csv", parse_dates=["date"])
    rpm = tt.posted_rate / tt.distance
    lane = tt.pickup + " > " + tt.delivery
    norm = rpm.groupby([lane, tt.equipment]).transform("median")
    ratio = rpm / norm
    quote_ratio = tt.quote_signal / norm
    mi_train = tt.groupby(tt.date.dt.month).market_index.mean()
    return {
        "train_rows": len(tt), "valid_rows": len(va),
        "train_dates": f"{tt.date.min():%d %b} to {tt.date.max():%d %b %Y}",
        "valid_dates": f"{va.date.min():%d %b} to {va.date.max():%d %b %Y}",
        "neg_weight": int((tt.weight < 0).sum()), "missing_weight": int(tt.weight.isna().sum()),
        "missing_market": int(tt.market_index.isna().sum()),
        "within7": float(((ratio - 1).abs() <= 0.07).mean() * 100),
        "rate_errors": float((~ratio.between(0.6, 1.7)).mean() * 100),
        "quote_off": float(((quote_ratio - 1).abs() > 0.25).mean() * 100),
        "mi_peak": float(mi_train.max()), "mi_peak_month": pd.Timestamp(2025, int(mi_train.idxmax()), 1).strftime("%B"),
        "mi_sep_oct": float(mi_train.loc[[9, 10]].mean()),
        "mi_valid": float(va.market_index.mean()),
    }


def backtest_chart(table):
    keep = ["lane median rate/mile x distance", "LightGBM + market_index + quote_signal", "LightGBM + market_index",
            "LightGBM, no market signals, recency-weighted (final)"]
    labels = ["Lane median", "LGBM + market + quote", "LGBM + market", "Final model"]
    pivot = table[table.model.isin(keep)].pivot(index="model", columns="window", values="MAPE clean %").loc[keep]
    fig, ax = plt.subplots(figsize=(7.5, 3.4), dpi=160)
    x = range(len(keep))
    w = 0.38
    for i, (col, color) in enumerate(zip(pivot.columns, ["#7FA7AE", "#064A56"])):
        ax.bar([v + (i - 0.5) * w for v in x], pivot[col], width=w, label=col, color=color)
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontsize=9)
    ax.set_ylabel("MAPE on normal loads (%)")
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=9)
    fig.tight_layout()
    path = REPORTS / "backtest_chart.png"
    fig.savefig(path)
    plt.close(fig)
    return path


def add_table(doc, frame):
    t = doc.add_table(rows=1, cols=len(frame.columns))
    t.style = "Light Grid Accent 1"
    for cell, name in zip(t.rows[0].cells, frame.columns):
        cell.text = str(name)
    for _, row in frame.iterrows():
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = f"{value:.2f}" if isinstance(value, float) else str(value)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)


def main():
    table = pd.read_csv(REPORTS / "backtest_results.csv")
    d = data_checks()
    chart = backtest_chart(table)
    final = table[table.model.str.contains("final")].set_index("window")
    lane = table[table.model.str.startswith("lane median")].set_index("window")
    dec = pd.read_csv(ROOT / "data" / "december_chart_inputs.csv")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    doc.add_heading("Freight Rate Prediction", level=0)
    doc.add_paragraph("Masood Ahmad · Spotter ML Engineer assessment")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(
        "The final model prices a load from its lane, equipment, distance, weight and day, using LightGBM on log rate "
        "per mile with an L1 loss, recency weighting and a three-seed average. It leaves out quote_signal and "
        "market_index on purpose: neither held up when tested forward in time. "
        f"On the most recent backtest (train Jan to Aug, predict Sep to Oct) the model's error is {final.loc['B: Sep-Oct', 'MAPE clean %']:.1f}% "
        f"on normal loads and {final.loc['B: Sep-Oct', 'MAPE all %']:.1f}% on all loads, against "
        f"{lane.loc['B: Sep-Oct', 'MAPE clean %']:.1f}% and {lane.loc['B: Sep-Oct', 'MAPE all %']:.1f}% for a lane-average baseline.")

    doc.add_heading("Data", level=1)
    doc.add_paragraph(
        f"Training covers {d['train_rows']:,} loads from {d['train_dates']}. Validation is {d['valid_rows']:,} loads from "
        f"{d['valid_dates']}, so the task is a two-month forecast past the training data. Issues found and how I handled them:")
    for line in [
        f"{d['neg_weight']} negative weights. The magnitudes are in the normal 15k to 47k lb range, so I treat them as sign flips and take the absolute value, keeping a flag.",
        f"{d['missing_weight']} missing weights, filled with the median for that equipment type, with a flag.",
        f"{d['missing_market']} missing market_index values, filled with that day's median.",
        f"Rates are tightly clustered: {d['within7']:.0f}% of loads sit within 7% of the median rate per mile for their lane and equipment.",
        f"About {d['rate_errors']:.1f}% of rates are 0.2 to 0.4x or 2.5 to 4.5x their lane norm. These look like unit or entry errors. I drop them from training only; they stay in every evaluation number marked 'all'.",
        f"quote_signal is off its lane norm by more than 25% on {d['quote_off']:.0f}% of loads, so it behaves like a noisy hint rather than a price.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Split approach", level=1)
    doc.add_paragraph(
        "Random splits would leak future prices into training, so every evaluation is forward in time with the same "
        "two-month horizon as the real task. Window A trains on January to June and predicts July and August. "
        "Window B trains on January to August and predicts September and October. Results are reported on all test "
        "loads and on normal loads (excluding the rate errors described above, which no model can predict). "
        "The final model is then retrained on all of January to October before predicting November and December.")

    doc.add_heading("Validation results", level=1)
    show = table[["window", "model", "MAPE all %", "MAE all $", "MAPE clean %", "median APE clean %"]]
    add_table(doc, show)
    doc.add_paragraph()
    doc.add_picture(str(chart), width=Inches(6.2))
    doc.add_paragraph(
        "Two things decided the model. First, anchoring on distance x quote_signal is the worst option because the quote "
        "is wrong on a sizeable share of loads. Second, market_index is unstable: it helped slightly in window A and "
        "hurt badly in window B, where the index dropped to "
        f"{d['mi_sep_oct']:.2f} after peaking at {d['mi_peak']:.2f} in {d['mi_peak_month']} while actual rates held up. "
        f"The validation months sit in that same low range ({d['mi_valid']:.2f}), so the model without it is the safer "
        "choice for the period being scored. Among the models without market signals, a 90-day recency weight was better in both windows, most clearly in the later one.")

    doc.add_heading("December forecast", level=1)
    doc.add_paragraph(
        "Lexington to Fort Wayne, 360 miles, Dry Van, 32,000 lb, one prediction per day. The December inputs carry no "
        "quote_signal or market_index, which the final model does not need. City coordinates come from the same cities "
        f"in the other files. Predicted rates run from ${dec.predicted_rate.min():,.2f} to ${dec.predicted_rate.max():,.2f}, "
        "with a weekly pattern (midweek higher, weekends lower) that matches the training data. Recent dry-van loads on "
        "this lane paid about $780 to $865 for similar distances, so the level is consistent with history.")
    doc.add_picture(str(ROOT / "scorer_results" / "candidate_december.png"), width=Inches(6.4))

    doc.add_heading("Limits and next steps", level=1)
    for line in [
        "There is no previous December in the data, so holiday effects around Christmas are not modelled.",
        "The rate errors in the target can't be predicted from the inputs. If they come from a known source (per-mile vs total, currency), fixing them upstream would be worth more than any model change.",
        "With more history, I would add a lane-level time trend and retest market_index with a lag instead of same-day.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    out = REPORTS / "Freight_Rate_Report.docx"
    doc.save(out)
    print(f"report written: {out}")


if __name__ == "__main__":
    main()
